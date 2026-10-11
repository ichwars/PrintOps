"""Bounded plain-text spreadsheet projections for the file manager."""

import csv
import io
import uuid
import zipfile
from itertools import islice
from pathlib import Path

from openpyxl import load_workbook

from backend.app.utils.archive_budget import ArchiveBudget, ArchiveBudgetError, read_xml_member, validate_zip_archive

MAX_PREVIEW_BYTES = 20 * 1024 * 1024
MAX_ROWS = 500
MAX_COLS = 40
MAX_SHEETS = 20
MAX_CELL_CHARS = 2000
MAX_TOTAL_CHARS = 300_000
MAX_TOTAL_CELLS = 50_000
_BUDGET = ArchiveBudget(
    max_members=1024,
    max_member_bytes=MAX_PREVIEW_BYTES,
    max_total_uncompressed_bytes=MAX_PREVIEW_BYTES,
    max_metadata_bytes=MAX_PREVIEW_BYTES,
)
_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def read_preview_bytes(path: Path) -> bytes:
    if path.stat().st_size > MAX_PREVIEW_BYTES:
        raise ArchiveBudgetError("Document exceeds the 20 MiB preview limit")
    with path.open("rb") as source:
        data = source.read(MAX_PREVIEW_BYTES + 1)
    if len(data) > MAX_PREVIEW_BYTES:
        raise ArchiveBudgetError("Document exceeds the 20 MiB preview limit")
    return data


def _text(value) -> str:
    return "" if value is None else str(value)[:MAX_CELL_CHARS]


def _sheet(name, rows, total_rows, total_cols):
    return {"name": _text(name), "rows": rows, "totalRows": total_rows, "totalCols": total_cols}


def spreadsheet_preview(path: Path, file_type: str) -> dict:
    data = read_preview_bytes(path)
    sheets = []
    truncated = False
    remaining_chars = MAX_TOTAL_CHARS
    remaining_cells = MAX_TOTAL_CELLS

    def can_project():
        return remaining_chars > 0 and remaining_cells > 0

    def project_row(values):
        nonlocal truncated, remaining_chars, remaining_cells
        if not can_project():
            truncated = True
            return None
        projected = []
        for value in islice(values, MAX_COLS):
            if not can_project():
                truncated = True
                break
            text = "" if value is None else str(value)
            limit = min(MAX_CELL_CHARS, remaining_chars)
            truncated |= len(text) > limit
            projected.append(text[:limit])
            remaining_chars -= len(projected[-1])
            remaining_cells -= 1
        return projected

    if file_type == "csv":
        text = data.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text, newline=""))
        rows = []
        total_rows = total_cols = 0
        for row in reader:
            total_rows += 1
            total_cols = max(total_cols, len(row))
            if len(rows) < MAX_ROWS:
                projected = project_row(row)
                if projected is not None:
                    rows.append(projected)
        sheets.append(_sheet(path.name, rows, total_rows, total_cols))
    elif file_type in ("xlsx", "ods"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = validate_zip_archive(archive, _BUDGET)
            if file_type == "xlsx":
                for info in infos:
                    if info.filename.endswith((".xml", ".rels")):
                        read_xml_member(archive, info.filename, budget=_BUDGET)
            if file_type == "xlsx":
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    truncated = len(workbook.worksheets) > MAX_SHEETS
                    for worksheet in workbook.worksheets[:MAX_SHEETS]:
                        rows = []
                        if can_project():
                            for row in worksheet.iter_rows(
                                min_row=1,
                                max_row=min(worksheet.max_row or MAX_ROWS, MAX_ROWS),
                                max_col=min(worksheet.max_column or MAX_COLS, MAX_COLS),
                                values_only=True,
                            ):
                                projected = project_row(row)
                                if projected is None:
                                    break
                                rows.append(projected)
                        sheets.append(
                            _sheet(worksheet.title, rows, worksheet.max_row or len(rows), worksheet.max_column or 0)
                        )
                finally:
                    workbook.close()
            else:
                root = read_xml_member(archive, "content.xml", budget=_BUDGET)
                tables = list(islice(root.iter(_TABLE + "table"), MAX_SHEETS + 1))
                truncated = len(tables) > MAX_SHEETS
                for table in tables[:MAX_SHEETS]:
                    rows = []
                    total_rows = total_cols = 0
                    for row in table.findall(_TABLE + "table-row"):
                        row_repeat = min(int(row.get(_TABLE + "number-rows-repeated", "1")), 1_048_576)
                        if row_repeat < 1:
                            raise ValueError("Invalid row repetition")
                        width = 0
                        cells = []
                        for cell in row:
                            if cell.tag not in (_TABLE + "table-cell", _TABLE + "covered-table-cell"):
                                continue
                            repeat = min(int(cell.get(_TABLE + "number-columns-repeated", "1")), 16_384)
                            if repeat < 1:
                                raise ValueError("Invalid column repetition")
                            cells.append((cell, repeat))
                            width += repeat
                        total_cols = max(total_cols, width)
                        total_rows += row_repeat
                        for _ in range(min(row_repeat, max(0, MAX_ROWS - len(rows)))):
                            if not can_project():
                                truncated = True
                                break

                            def values(source_cells=cells):
                                for cell, repeat in source_cells:
                                    value = "\n".join("".join(p.itertext()) for p in cell.findall(_TEXT + "p"))
                                    for _ in range(min(repeat, MAX_COLS)):
                                        yield value

                            projected = project_row(values())
                            if projected is not None:
                                rows.append(projected)
                    sheets.append(_sheet(table.get(_TABLE + "name", "Sheet"), rows, total_rows, total_cols))
    else:
        raise ValueError("Preview supports CSV, XLSX and ODS")
    return {
        "sheets": sheets,
        "truncated": truncated
        or any(sheet["totalRows"] > len(sheet["rows"]) or sheet["totalCols"] > MAX_COLS for sheet in sheets),
    }


def generate_pdf_thumbnail(path: Path, thumbnails_dir: Path) -> str | None:
    """Best-effort bounded first-page raster; callers run this off the event loop."""
    from backend.app.services.pdf_render import render_pdf_thumbnail

    try:
        content = render_pdf_thumbnail(read_preview_bytes(path))
        destination = thumbnails_dir / f"{uuid.uuid4().hex}.png"
        destination.write_bytes(content)
        return str(destination)
    except (ValueError, OSError):
        return None
